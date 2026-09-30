import re
import uuid
from pathlib import Path

from app.models import DocumentType


def detect_doc_type(filename: str) -> DocumentType:
    ext = Path(filename).suffix.lower()
    mapping = {
        ".pdf": DocumentType.PDF,
        ".pptx": DocumentType.PPTX,
        ".docx": DocumentType.DOCX,
        ".xlsx": DocumentType.XLSX,
        ".csv": DocumentType.CSV,
        ".txt": DocumentType.TXT,
        ".md": DocumentType.MD,
        ".html": DocumentType.HTML,
    }
    return mapping.get(ext, DocumentType.UNKNOWN)


def parse_pdf(file_path: str) -> tuple[list[dict], int]:
    import pdfplumber

    pages = []
    with pdfplumber.open(file_path) as pdf:
        for i, page in enumerate(pdf.pages):
            text = page.extract_text() or ""
            tables = page.extract_tables()
            table_text = ""
            for table in tables:
                if table:
                    for row in table:
                        if row:
                            table_text += " | ".join(str(c or "") for c in row) + "\n"
            pages.append(
                {
                    "page_num": i + 1,
                    "content": text,
                    "tables": table_text,
                    "full_text": text + "\n" + table_text,
                }
            )
        return pages, len(pdf.pages)


def parse_pptx(file_path: str) -> tuple[list[dict], int]:
    from pptx import Presentation

    prs = Presentation(file_path)
    pages = []
    for i, slide in enumerate(prs.slides):
        text_parts = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text_parts.append(shape.text)
        pages.append(
            {
                "page_num": i + 1,
                "content": "\n".join(text_parts),
                "tables": "",
                "full_text": "\n".join(text_parts),
            }
        )
    return pages, len(prs.slides)


def parse_docx(file_path: str) -> tuple[list[dict], int]:
    from docx import Document as DocxDocument

    doc = DocxDocument(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    text = "\n".join(paragraphs)
    pages = [{"page_num": 1, "content": text, "tables": "", "full_text": text}]
    return pages, 1


def parse_xlsx(file_path: str) -> tuple[list[dict], int]:
    from openpyxl import load_workbook

    wb = load_workbook(file_path, data_only=True)
    pages = []
    for i, sheet_name in enumerate(wb.sheetnames):
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append(" | ".join(str(c or "") for c in row))
        text = "\n".join(rows)
        pages.append({"page_num": i + 1, "content": text, "tables": text, "full_text": text})
    return pages, len(wb.sheetnames)


def parse_text(file_path: str) -> tuple[list[dict], int]:
    with open(file_path, encoding="utf-8", errors="ignore") as f:
        text = f.read()
    pages = [{"page_num": 1, "content": text, "tables": "", "full_text": text}]
    return pages, 1


PARSERS = {
    ".pdf": parse_pdf,
    ".pptx": parse_pptx,
    ".docx": parse_docx,
    ".xlsx": parse_xlsx,
    ".txt": parse_text,
    ".md": parse_text,
    ".html": parse_text,
    ".csv": parse_text,
}


def parse_document(file_path: str, filename: str) -> tuple[list[dict], int]:
    ext = Path(filename).suffix.lower()
    parser = PARSERS.get(ext)
    if parser:
        return parser(file_path)
    return parse_text(file_path)


def chunk_pages(pages: list[dict], chunk_size: int = 800, overlap: int = 200) -> list[dict]:
    chunks = []
    for page in pages:
        text = page["full_text"]
        if not text.strip():
            continue

        sentences = re.split(r"(?<=[.!?\n])\s+", text)
        current_chunk = []
        current_len = 0

        for sentence in sentences:
            sent_len = len(sentence.split())
            if current_len + sent_len > chunk_size and current_chunk:
                chunk_text = " ".join(current_chunk)
                chunks.append(
                    {
                        "id": str(uuid.uuid4()),
                        "text": chunk_text,
                        "page_num": page["page_num"],
                        "word_count": current_len,
                    }
                )
                overlap_words = int(overlap / 5)
                current_chunk = (
                    current_chunk[-overlap_words:] if len(current_chunk) > overlap_words else []
                )
                current_len = sum(len(s.split()) for s in current_chunk)

            current_chunk.append(sentence)
            current_len += sent_len

        if current_chunk:
            chunks.append(
                {
                    "id": str(uuid.uuid4()),
                    "text": " ".join(current_chunk),
                    "page_num": page["page_num"],
                    "word_count": current_len,
                }
            )

    return chunks
